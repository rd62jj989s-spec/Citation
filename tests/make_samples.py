"""Erzeugt erfundene Beispielquellen und eine Beispiel-Masterarbeit für die Tests.

Alle Texte, Autorinnen, Autoren und Verlage sind frei erfunden.

Ergebnis im Ordner samples/:
  Muster_2021_Organisationskommunikation.pdf  Buchauszug, PDF-Seite 1 Titelei, PDF-Seiten 2 bis 5 = S. 157 bis 160
  Beispiel_2019_Digitale_Oeffentlichkeiten.pdf  Zeitschriftenartikel, Seitenzahl oben rechts, S. 41 bis 44
  masterarbeit_beispiel.docx                   Beispielarbeit mit APA-7-Belegen
  masterarbeit_beispiel.pdf                    dieselbe Arbeit als PDF
  masterarbeit_beispiel_v2.docx                korrigierte Fassung zum Test von „Aktualisieren“
"""

import os

from docx import Document
from docx.shared import Cm, Pt
from reportlab.lib.pagesizes import A5
from reportlab.pdfgen import canvas

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "samples")

# Buchauszug: jede Zeile ist bewusst vorgebrochen, inklusive Silbentrennung.
BOOK_PAGES = {
    157: {
        "heading": ["6 Vertrauen als Ressource der", "Organisationskommunikation"],
        "lines": [
            "Organisationen kommunizieren nicht im luftleeren Raum. Sie",
            "sind auf das Vertrauen ihrer Anspruchsgruppen angewiesen,",
            "das sich nur über längere Zeiträume aufbauen lässt. Ver-",
            "trauen ist dabei kein Zustand, der einmal erreicht und dann",
            "dauerhaft gesichert wäre, sondern ein fortlaufender Prozess",
            "der wechselseitigen Beobachtung.",
            "\tAus Sicht der strategischen Kommunikation lässt sich Ver-",
            "trauen als Erwartung beschreiben, dass sich eine Organisa-",
            "tion auch in unübersichtlichen Situationen verlässlich ver-",
            "hält. Diese Erwartung entsteht aus früheren Erfahrungen,",
            "aus der wahrgenommenen Kompetenz der Organisation und",
            "aus der Übereinstimmung zwischen öffentlichen Aussagen",
            "und tatsächlichem Handeln.",
        ],
    },
    158: {
        "lines": [
            "Glaubwürdigkeit entsteht dort, wo Worte und Taten einer",
            "Organisation dauerhaft übereinstimmen. Wird diese Über-",
            "einstimmung auch nur punktuell verletzt, verlieren Anspruchs-",
            "gruppen rasch das Vertrauen, das über Jahre aufgebaut wurde.",
            "\tFür die Praxis folgt daraus, dass Kommunikationsabteilun-",
            "gen nicht allein für die Gestaltung von Botschaften zuständig",
            "sind. Sie müssen vielmehr frühzeitig in Entscheidungs-",
            "prozesse eingebunden werden, um die kommunikativen Folgen",
            "von Managemententscheidungen abschätzen zu können.",
            "\tDie Bedeutung dieser Einbindung wird häufig unterschätzt.",
            "Gerade in Unternehmen mit flachen Hierarchien entscheidet",
            "sich früh, welche Themen später öffentlich verhandelt werden.",
        ],
    },
    159: {
        "lines": [
            "Die interne Kommunikation bildet das Fundament jeder glaub-",
            "würdigen Außendarstellung. Mitarbeitende, die über Ziele",
            "und Hintergründe informiert sind, tragen die Botschaften",
            "ihrer Organisation mit und wirken als ihre Botschafterinnen",
            "und Botschafter.¹",
            "\tIn Krisensituationen zeigt sich, ob eine Organisation zuvor",
            "in belastbare Beziehungen investiert hat. Wer erst im Ernst-",
            "fall beginnt, mit seinen Anspruchsgruppen zu sprechen, hat",
            "bereits einen entscheidenden Teil seiner Handlungsfähigkeit",
        ],
        "footnote": "¹ Vgl. dazu ausführlich Kapitel 4 in diesem Band.",
    },
    160: {
        "lines": [
            "verloren. Krisenkommunikation ist deshalb weniger eine",
            "Frage der Technik als eine Frage der Beziehungen, die vor",
            "der Krise gepflegt wurden.",
            "\tDaraus ergibt sich eine doppelte Aufgabe für die Kommu-",
            "nikationsverantwortlichen: Sie müssen Beziehungen im All-",
            "tag pflegen und zugleich Strukturen schaffen, die im Ernst-",
            "fall schnelle Entscheidungen ermöglichen.",
        ],
    },
}


def make_book(path):
    w, h = A5
    c = canvas.Canvas(path, pagesize=A5)
    c.setTitle("Grundlagen der Organisationskommunikation")
    c.setAuthor("Anna Muster")
    # Titelei ohne Seitenzahl
    c.setFont("Times-Bold", 16)
    c.drawCentredString(w / 2, h - 170, "Anna Muster")
    c.setFont("Times-Bold", 18)
    c.drawCentredString(w / 2, h - 210, "Grundlagen der")
    c.drawCentredString(w / 2, h - 234, "Organisationskommunikation")
    c.setFont("Times-Roman", 11)
    c.drawCentredString(w / 2, h - 290, "Beispielverlag 2021")
    c.drawCentredString(w / 2, 90, "Fiktive Beispielquelle für den Zitationsprüfer")
    c.showPage()
    for num, page in BOOK_PAGES.items():
        c.setFont("Times-Roman", 8.5)
        c.drawCentredString(w / 2, h - 32, "Vertrauen als Ressource")
        y = h - 70
        if "heading" in page:
            c.setFont("Times-Bold", 12)
            for line in page["heading"]:
                c.drawString(50, y, line)
                y -= 16
            y -= 10
        c.setFont("Times-Roman", 10.5)
        for line in page["lines"]:
            x = 50
            if line.startswith("\t"):
                line = line[1:]
                x = 62
            c.drawString(x, y, line)
            y -= 14
        if "footnote" in page:
            c.setFont("Times-Roman", 8)
            c.drawString(50, 62, page["footnote"])
        c.setFont("Times-Roman", 9)
        c.drawCentredString(w / 2, 30, str(num))
        c.showPage()
    c.save()


ARTICLE_PAGES = [
    (
        "Digitale Öffentlichkeiten und organisationale Sichtbarkeit",
        [
            "Plattformen haben die Bedingungen, unter denen Organisationen öffentlich sichtbar werden, "
            "grundlegend verschoben. Sichtbarkeit entsteht nicht mehr vorrangig über redaktionelle Auswahl, "
            "sondern über Empfehlungslogiken, die sich an Interaktionen orientieren.",
            "Für Kommunikationsabteilungen bedeutet das, dass sie ihre Reichweite nur begrenzt planen können. "
            "Sie konkurrieren mit privaten Beiträgen, Werbung und Nachrichten um dieselbe Aufmerksamkeit.",
        ],
    ),
    (
        None,
        [
            "Die Beziehungspflege verlagert sich damit in Räume, die Organisationen nicht selbst kontrollieren. "
            "Wer auf Plattformen kommuniziert, gibt einen Teil der Deutungshoheit an Algorithmen und an das "
            "Publikum ab.",
            "Gleichzeitig eröffnen digitale Öffentlichkeiten die Möglichkeit, unmittelbar mit Anspruchsgruppen "
            "in einen Dialog zu treten, ohne den Umweg über journalistische Medien zu gehen.",
        ],
    ),
    (
        None,
        [
            "Dialogorientierte Kommunikation verlangt allerdings Ressourcen, die in vielen Organisationen fehlen. "
            "Schnelle Reaktionszeiten und ein verbindlicher Ton lassen sich nur mit klaren Zuständigkeiten "
            "gewährleisten.",
            "Studien zeigen, dass Organisationen mit festen Redaktionsteams auf Plattformen deutlich häufiger "
            "als glaubwürdig wahrgenommen werden als Organisationen, die nur anlassbezogen kommunizieren.",
        ],
    ),
    (
        None,
        [
            "Insgesamt ist digitale Beziehungspflege weniger eine technische als eine organisatorische Aufgabe. "
            "Sie gelingt dort, wo Kommunikation als dauerhafte Aufgabe der gesamten Organisation verstanden wird.",
        ],
    ),
]


def make_article_pdf(path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import simpleSplit

    w, h = A4
    c = canvas.Canvas(path, pagesize=A4)
    c.setTitle("Digitale Öffentlichkeiten und organisationale Sichtbarkeit")
    c.setAuthor("Mira Beispiel")
    for i, (title, paras) in enumerate(ARTICLE_PAGES):
        num = 41 + i
        c.setFont("Helvetica", 8)
        c.drawString(60, h - 40, "Zeitschrift für Beispielforschung, 7(2)")
        c.drawRightString(w - 60, h - 40, str(num))
        y = h - 90
        if title:
            c.setFont("Helvetica", 11)
            c.drawString(60, y, "Mira Beispiel")
            y -= 24
            c.setFont("Helvetica-Bold", 15)
            c.drawString(60, y, title)
            y -= 34
        c.setFont("Helvetica", 11)
        for text in paras:
            for line in simpleSplit(text, "Helvetica", 11, w - 120):
                c.drawString(60, y, line)
                y -= 15
            y -= 9
        c.showPage()
    c.save()


def make_thesis_pdf(path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import simpleSplit

    w, h = A4
    c = canvas.Canvas(path, pagesize=A4)
    page = 1
    y = h - 80

    def new_page():
        nonlocal y, page
        c.setFont("Helvetica", 9)
        c.drawCentredString(w / 2, 40, str(page))
        c.showPage()
        page += 1
        y = h - 80

    for kind, text in THESIS:
        if kind == "h":
            font, size, x, indent = "Helvetica-Bold", 12, 70, 0
            y -= 8
        elif kind == "q":
            font, size, x, indent = "Helvetica", 11, 106, 0
        elif kind == "b":
            font, size, x, indent = "Helvetica", 11, 70, 36
        else:
            font, size, x, indent = "Helvetica", 11, 70, 0
        width = w - 70 - x
        lines = simpleSplit(text, font, size, width - indent)
        for k, line in enumerate(lines):
            if y < 90:
                new_page()
            c.setFont(font, size)
            c.drawString(x + (indent if k > 0 else 0), y, line)
            y -= 16
        y -= 8
    new_page()
    c.save()




THESIS = [
    ("h", "1 Einleitung"),
    ("p", "Diese Beispielarbeit dient nur dazu, den Zitationsprüfer zu testen. Alle Quellen sind erfunden."),
    ("h", "2 Vertrauen in der Organisationskommunikation"),
    ("p", "Vertrauen gilt als zentrale Ressource. Muster (2021) beschreibt Vertrauen als „fortlaufender Prozess der wechselseitigen Beobachtung“ (S. 157)."),
    ("p", "Besonders deutlich wird dies an der Glaubwürdigkeit: „Glaubwürdigkeit entsteht dort, wo Worte und Taten einer Organisation dauerhaft übereinstimmen“ (Muster, 2021, S. 158)."),
    ("p", "Muster (2021) fasst den Begriff genauer:"),
    ("q", "Aus Sicht der strategischen Kommunikation lässt sich Vertrauen als Erwartung beschreiben, dass sich eine Organisation auch in unübersichtlichen Situationen verlässlich verhält. Diese Erwartung entsteht aus früheren Erfahrungen, aus der wahrgenommenen Kompetenz der Organisation und aus der Übereinstimmung zwischen öffentlichen Aussagen und tatsächlichem Handeln. (S. 157)"),
    ("p", "Kommunikationsabteilungen müssen daher „frühzeitig in Entscheidungsprozesse eingebunden werden“ (Muster, 2021, S. 157)."),
    ("p", "Die interne Kommunikation gilt als „Fundament jeder glaubhaften Außendarstellung“ (Muster, 2021, S. 159)."),
    ("p", "„Wer erst im Ernstfall beginnt, mit seinen Anspruchsgruppen zu sprechen, hat bereits einen entscheidenden Teil seiner Handlungsfähigkeit verloren“ (Muster, 2021, S. 159)."),
    ("p", "„Mitarbeitende, die über Ziele und Hintergründe informiert sind, […] wirken als ihre Botschafterinnen und Botschafter“ (Muster, 2021, S. 159)."),
    ("p", "Krisenkommunikation hängt demnach vor allem von den zuvor gepflegten Beziehungen ab (Muster, 2021, S. 160)."),
    ("p", "Nach Muster (2021) ist „Vertrauen die wichtigste Währung jeder Organisation“ (S. 158)."),
    ("p", "Die Einbindung in Entscheidungen wird oft unterschätzt (vgl. Muster, 2021, S. 158 f.). Auch flache Hierarchien spielen eine Rolle (Muster 2021, S. 158)."),
    ("h", "3 Digitale Öffentlichkeiten"),
    ("p", "Auf Plattformen entsteht Sichtbarkeit „über Empfehlungslogiken, die sich an Interaktionen orientieren“ (Beispiel, 2019, S. 41)."),
    ("p", "Beispiel (2019, S. 42) betont, wer auf Plattformen kommuniziere, gebe „einen Teil der Deutungshoheit an Algorithmen und an das Publikum ab“."),
    ("p", "Organisationen mit festen Redaktionsteams werden als glaubwürdiger wahrgenommen (Beispiel, 2019, S. 43; Nichtda, 2020, S. 5)."),
    ("p", "Digitale Beziehungspflege ist vor allem eine organisatorische Aufgabe (Beispiel, 2019)."),
    ("h", "Literaturverzeichnis"),
    ("b", "Beispiel, M. (2019). Digitale Öffentlichkeiten und organisationale Sichtbarkeit. Zeitschrift für Beispielforschung, 7(2), 41–44."),
    ("b", "Muster, A. (2021). Grundlagen der Organisationskommunikation. Beispielverlag."),
    ("b", "Ungenutzt, U. (2018). Ein Buch, das nie zitiert wird. Musterverlag."),
]


# Korrigierte Fassung: Wortlaut und zwei Seitenangaben berichtigt
THESIS_V2 = [
    (
        kind,
        text.replace("„frühzeitig in Entscheidungsprozesse eingebunden werden“ (Muster, 2021, S. 157)", "„frühzeitig in Entscheidungsprozesse eingebunden werden“ (Muster, 2021, S. 158)")
        .replace("glaubhaften Außendarstellung", "glaubwürdigen Außendarstellung")
        .replace("Handlungsfähigkeit verloren“ (Muster, 2021, S. 159)", "Handlungsfähigkeit verloren“ (Muster, 2021, S. 159–160)"),
    )
    for kind, text in THESIS
]


def make_thesis_docx(path, content=THESIS):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Liberation Sans"
    style.font.size = Pt(11)
    for kind, text in content:
        if kind == "h":
            doc.add_heading(text, level=1)
        elif kind == "q":
            p = doc.add_paragraph(text)
            p.paragraph_format.left_indent = Cm(1.27)
        elif kind == "b":
            p = doc.add_paragraph(text)
            p.paragraph_format.left_indent = Cm(1.27)
            p.paragraph_format.first_line_indent = Cm(-1.27)
        else:
            doc.add_paragraph(text)
    doc.save(path)



def main():
    os.makedirs(OUT, exist_ok=True)
    make_book(os.path.join(OUT, "Muster_2021_Organisationskommunikation.pdf"))
    make_article_pdf(os.path.join(OUT, "Beispiel_2019_Digitale_Oeffentlichkeiten.pdf"))
    make_thesis_docx(os.path.join(OUT, "masterarbeit_beispiel.docx"))
    make_thesis_pdf(os.path.join(OUT, "masterarbeit_beispiel.pdf"))
    make_thesis_docx(os.path.join(OUT, "masterarbeit_beispiel_v2.docx"), THESIS_V2)
    print("Beispieldateien erstellt in", OUT)


if __name__ == "__main__":
    main()
